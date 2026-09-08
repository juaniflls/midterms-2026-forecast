from __future__ import annotations
import hashlib, sys
from pathlib import Path
import nbformat
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
required = [
    "Modelo_Midterms_2026_v27.1_NUCLEUS42_FINAL.ipynb",
    "Election_Model_v27_1_Coherence_Audit.html",
    "Model.xlsx",
    "outputs/Election_Model_Final_Report_v27_1.xlsx",
    "outputs/Model_Sensitivity_Audit_v27.xlsx",
    "outputs/scenario_state_engine_v27.py",
    "dash_app/app.py",
    "README.md",
    "MODEL_CARD.md",
    "CITATION.cff",
    "LICENSE",
]
fail=[]
for rel in required:
    if not (ROOT/rel).exists():
        fail.append(f"missing {rel}")
for junk in [".DS_Store","__MACOSX",".ipynb_checkpoints","__pycache__"]:
    if list(ROOT.rglob(junk)):
        fail.append(f"packaging junk: {junk}")

if not fail:
    nb=nbformat.read(ROOT/"Modelo_Midterms_2026_v27.1_NUCLEUS42_FINAL.ipynb", as_version=4)
    nbformat.validate(nb)
    code=[c for c in nb.cells if c.cell_type=="code"]
    errs=[o for c in code for o in c.get("outputs",[]) if o.get("output_type")=="error"]
    if len(code)!=13: fail.append(f"expected 13 code cells; found {len(code)}")
    if errs: fail.append(f"stored notebook errors: {len(errs)}")

    html=(ROOT/"Election_Model_v27_1_Coherence_Audit.html").read_text(encoding="utf-8")
    html_upper = html.upper()
    for tok in ["NUCLEUS 42","D 230","R 205","D 50","R 50","50,000 MONTE CARLO"]:
        if tok.upper() not in html_upper: fail.append(f"HTML missing {tok}")

    wb=load_workbook(ROOT/"outputs/Election_Model_Final_Report_v27_1.xlsx", read_only=True, data_only=True)
    needed={"ExecutiveSummary","FinalProjection","Dashboard_Data","HouseRaceDetail",
            "SenateRaceDetail","SenateSeatAccounting","CentralForecastContract",
            "ConsistencyAudit","ScenarioEngineContract"}
    miss=needed-set(wb.sheetnames)
    if miss: fail.append("report missing sheets: "+", ".join(sorted(miss)))
    for sheet in ["ConsistencyAudit","CentralForecastContract"]:
        if sheet in wb.sheetnames:
            rows=list(wb[sheet].iter_rows(values_only=True))
            if len(rows)>1:
                hdr=[str(x) if x is not None else "" for x in rows[0]]
                if "Passed" not in hdr:
                    fail.append(f"{sheet} lacks Passed")
                else:
                    j=hdr.index("Passed")
                    vals=[bool(r[j]) for r in rows[1:] if len(r)>j and r[j] is not None]
                    if vals and not all(vals): fail.append(f"{sheet} contains FAIL")
    wb.close()

if fail:
    print("STATUS: FAIL")
    for x in fail: print("[FAIL]",x)
    sys.exit(1)

print("STATUS: OK")
print("NUCLEUS 42 v27.1.0")
print("Notebook: 13 code cells · no stored errors")
print("House: D230-R205")
print("Senate: D50-R50")
